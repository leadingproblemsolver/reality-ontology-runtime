from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Any

from . import core


_REQUIRED_BY_ROLE = {
    "activities": {"activity_id", "activity_name", "wbs", "planned_start", "planned_finish", "status"},
    "previous_activities": {"activity_id", "activity_name", "wbs", "planned_start", "planned_finish", "status"},
    "relationships": {"predecessor_id", "successor_id"},
    "requirements": {"scope_type", "scope_id", "requirement"},
}


def _canonical_headers(headers: list[str]) -> dict[str, str]:
    normalized = {core._norm(header): header for header in headers if header is not None}
    resolved: dict[str, str] = {}
    for canonical, aliases in core.ALIASES.items():
        for alias in aliases:
            if core._norm(alias) in normalized:
                resolved[canonical] = normalized[core._norm(alias)]
                break
    return resolved


def _sheet_report(name: str, headers: list[str], *, selected: bool, required: set[str]) -> dict[str, Any]:
    mappings = _canonical_headers(headers)
    mapped_headers = set(mappings.values())
    return {
        "name": name,
        "selected": selected,
        "discovered_headers": headers,
        "canonical_mappings": mappings,
        "missing_required_fields": sorted(required - set(mappings)),
        "unsupported_fields": [header for header in headers if header and header not in mapped_headers],
    }


def inspect_source(path: str | Path, role: str) -> dict[str, Any]:
    if role not in _REQUIRED_BY_ROLE:
        raise ValueError(f"unsupported source role: {role}")

    source = Path(path)
    if not source.exists():
        raise core.ConstructionInputError(f"source does not exist: {source}")

    required = _REQUIRED_BY_ROLE[role]
    suffix = source.suffix.lower()
    base: dict[str, Any] = {
        "role": role,
        "path": str(source),
        "format": suffix.lstrip("."),
        "sha256": core._digest(source),
        "required_fields": sorted(required),
    }

    if suffix == ".csv":
        with source.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            headers = [str(value or "").strip() for value in next(reader, [])]
        selected = _sheet_report("CSV", headers, selected=True, required=required)
        return {
            **base,
            "worksheet_names": [],
            "selected_worksheet": None,
            "worksheets": [selected],
            "canonical_mappings": selected["canonical_mappings"],
            "missing_required_fields": selected["missing_required_fields"],
            "unsupported_fields": selected["unsupported_fields"],
            "status": "PASS" if not selected["missing_required_fields"] else "FAIL",
        }

    if suffix in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise core.ConstructionInputError(
                "XLSX input requires openpyxl; install with `pip install .[xlsx]`"
            ) from exc

        workbook = load_workbook(source, read_only=True, data_only=True)
        names = list(workbook.sheetnames)
        selected_name = names[0] if names else None
        sheets: list[dict[str, Any]] = []
        for name in names:
            worksheet = workbook[name]
            iterator = worksheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(iterator, [])]
            sheets.append(_sheet_report(name, headers, selected=name == selected_name, required=required))
        selected = next((sheet for sheet in sheets if sheet["selected"]), None)
        return {
            **base,
            "worksheet_names": names,
            "selected_worksheet": selected_name,
            "worksheets": sheets,
            "canonical_mappings": selected["canonical_mappings"] if selected else {},
            "missing_required_fields": selected["missing_required_fields"] if selected else sorted(required),
            "unsupported_fields": selected["unsupported_fields"] if selected else [],
            "status": "PASS" if selected and not selected["missing_required_fields"] else "FAIL",
        }

    raise core.ConstructionInputError(f"unsupported source type: {suffix}; use CSV or XLSX")


def preflight_real_cycle(
    *,
    activities_path: str,
    relationships_path: str,
    requirements_path: str,
    previous_activities_path: str | None,
    output_dir: str,
) -> tuple[dict[str, Any], Path]:
    target = Path(output_dir)
    target.mkdir(parents=True, exist_ok=True)
    report_path = target / "input_preflight.json"

    sources = {
        "activities": activities_path,
        "relationships": relationships_path,
        "requirements": requirements_path,
    }
    if previous_activities_path:
        sources["previous_activities"] = previous_activities_path

    reports: dict[str, Any] = {}
    failures: list[str] = []
    for role, path in sources.items():
        try:
            report = inspect_source(path, role)
        except core.ConstructionInputError as exc:
            report = {
                "role": role,
                "path": str(path),
                "status": "FAIL",
                "error": str(exc),
            }
        reports[role] = report
        if report.get("status") != "PASS":
            missing = report.get("missing_required_fields") or []
            detail = f"missing {', '.join(missing)}" if missing else str(report.get("error", "preflight failed"))
            failures.append(f"{role}: {detail}")

    payload = {
        "contract": "construction_real_cycle_input_preflight_v1",
        "sources": reports,
        "status": "PASS" if not failures else "FAIL",
    }
    report_path.write_text(json.dumps(payload, sort_keys=True, indent=2) + "\n", encoding="utf-8")

    if failures:
        raise core.ConstructionInputError(
            f"real-input preflight failed; inspect {report_path}: " + "; ".join(failures)
        )
    return payload, report_path


def _write_engineer_review(machine_path: Path, review_path: Path) -> None:
    with machine_path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        fieldnames = list(reader.fieldnames or [])
        rows = list(reader)

    review_fields = ["engineer_verdict", "engineer_correction", "engineer_note"]
    with review_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames + review_fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({**row, **{field: "" for field in review_fields}})


def _recompute_receipt_digest(receipt: dict[str, Any]) -> str:
    payload = dict(receipt)
    payload.pop("projection_digest", None)
    text = json.dumps(payload, sort_keys=True, indent=2)
    return hashlib.sha256(text.encode()).hexdigest()


def run_handoff_ready_lookahead(
    *,
    activities_path: str,
    relationships_path: str,
    requirements_path: str,
    as_of,
    days: int = 90,
    output_dir: str = "artifacts/construction",
    previous_activities_path: str | None = None,
    with_impact: bool = False,
) -> dict[str, Any]:
    _, preflight_path = preflight_real_cycle(
        activities_path=activities_path,
        relationships_path=relationships_path,
        requirements_path=requirements_path,
        previous_activities_path=previous_activities_path,
        output_dir=output_dir,
    )

    receipt = core.run_lookahead(
        activities_path=activities_path,
        relationships_path=relationships_path,
        requirements_path=requirements_path,
        previous_activities_path=previous_activities_path,
        with_impact=with_impact,
        as_of=as_of,
        days=days,
        output_dir=output_dir,
    )

    target = Path(output_dir)
    review_path: Path | None = None
    impact_path = target / "change_impact.csv"
    if with_impact:
        review_path = target / "engineer_review.csv"
        _write_engineer_review(impact_path, review_path)

    source_paths = {
        "activities": activities_path,
        "relationships": relationships_path,
        "requirements": requirements_path,
    }
    if previous_activities_path:
        source_paths["previous_activities"] = previous_activities_path

    receipt["source_identity"] = {
        role: {"path": str(Path(path)), "sha256": core._digest(Path(path))}
        for role, path in source_paths.items()
    }
    receipt["outputs"]["preflight_json"] = str(preflight_path)
    if review_path is not None:
        receipt["outputs"]["engineer_review_csv"] = str(review_path)
    receipt["projection_digest"] = _recompute_receipt_digest(receipt)

    receipt_path = target / "projection_receipt.json"
    receipt_path.write_text(json.dumps(receipt, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return receipt
